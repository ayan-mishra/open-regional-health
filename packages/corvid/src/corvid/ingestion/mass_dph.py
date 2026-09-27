from dataclasses import dataclass
from datetime import date, datetime
from io import BytesIO

from openpyxl import load_workbook

from corvid.domain import GeographyRecord, ObservationRecord, SourceRecord
from corvid.ingestion.base import FetchResult, SourceSchemaError
from corvid.ingestion.client import SourceHttpClient
from corvid.normalization.geography import MASSACHUSETTS_GEOGRAPHIES

XLSX_MEDIA_TYPE = (
    "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
)


@dataclass(frozen=True, slots=True)
class DphWorkbookSignature:
    season: str
    sheet_names: tuple[str, ...]
    required_headers_by_sheet: dict[str, tuple[str, ...]]
    date_column_by_sheet: dict[str, str]
    metric_columns_by_sheet: dict[str, tuple[str, ...]]
    expected_date_type_by_sheet: dict[str, str]


@dataclass(frozen=True, slots=True)
class DphParsedRow:
    season: str
    sheet_name: str
    source_label: str
    original_date_header: str
    original_metric_header: str
    date_cell_value: object
    raw_value_text: str | None


class MassDphConnector:


    source_key = "mass_dph_archive"
    parser_version = "mass-dph-guard-1"
    discovery_url = "https://www.mass.gov/info-details/influenza-reporting"
    allowed_origins = frozenset({"https://www.mass.gov"})
    source = SourceRecord(
        key=source_key,
        name="Massachusetts DPH Influenza Archive",
        organization="Massachusetts Department of Public Health",
        landing_page_url=discovery_url,
        update_cadence="seasonal archive",
        allowed_use="research; commercial use requires review",
        revision_policy="manual workbook signature required per season",
    )
    geographies: tuple[GeographyRecord, ...] = (MASSACHUSETTS_GEOGRAPHIES[0],)

    def __init__(
        self,
        *,
        signature: DphWorkbookSignature | None = None,
        source_label: str = "2025-2026 archived influenza dashboard data",
    ) -> None:
        self.signature = signature
        self.source_label = source_label

    def fetch(self, client: SourceHttpClient) -> FetchResult:
        raise SourceSchemaError(
            "Mass. DPH discovery is manual until an official workbook URL is captured"
        )

    def parse(self, snapshot: bytes) -> list[DphParsedRow]:
        return self.parse_workbook(snapshot, content_type=XLSX_MEDIA_TYPE)

    def parse_workbook(
        self, content: bytes, *, content_type: str
    ) -> list[DphParsedRow]:
        if self.signature is None:
            raise SourceSchemaError(
                "a manually reviewed workbook signature is required"
            )
        if content_type.split(";", 1)[0].strip().lower() != XLSX_MEDIA_TYPE:
            raise SourceSchemaError("Mass. DPH response is not an XLSX media type")
        if not content.startswith(b"PK"):
            raise SourceSchemaError("Mass. DPH response is not an XLSX ZIP payload")
        try:
            workbook = load_workbook(BytesIO(content), read_only=True, data_only=True)
        except (OSError, ValueError, KeyError) as error:
            raise SourceSchemaError("Mass. DPH XLSX cannot be inspected") from error
        if tuple(workbook.sheetnames) != self.signature.sheet_names:
            raise SourceSchemaError("Mass. DPH workbook has unexpected sheet names")

        parsed: list[DphParsedRow] = []
        for sheet_name in self.signature.sheet_names:
            sheet = workbook[sheet_name]
            values = sheet.iter_rows(values_only=True)
            try:
                raw_headers = next(values)
            except StopIteration as error:
                raise SourceSchemaError(
                    f"sheet {sheet_name} has no header row"
                ) from error
            headers = _header_map(raw_headers, sheet_name)
            required = self.signature.required_headers_by_sheet[sheet_name]
            for header in required:
                if _normalize_header(header) not in headers:
                    raise SourceSchemaError(
                        f"sheet {sheet_name} is missing required header {header}"
                    )
            date_header = self.signature.date_column_by_sheet[sheet_name]
            date_index, original_date_header = headers[_normalize_header(date_header)]
            metric_headers = self.signature.metric_columns_by_sheet[sheet_name]
            metric_columns = [
                headers[_normalize_header(header)] for header in metric_headers
            ]
            expected_type = self.signature.expected_date_type_by_sheet[sheet_name]
            for row_number, row in enumerate(values, start=2):
                date_value = row[date_index] if date_index < len(row) else None
                if date_value is None:
                    continue
                if expected_type == "date" and not isinstance(
                    date_value, (date, datetime)
                ):
                    raise SourceSchemaError(
                        f"sheet {sheet_name} row {row_number} date type changed"
                    )
                for metric_index, original_metric_header in metric_columns:
                    value = row[metric_index] if metric_index < len(row) else None
                    parsed.append(
                        DphParsedRow(
                            season=self.signature.season,
                            sheet_name=sheet_name,
                            source_label=self.source_label,
                            original_date_header=original_date_header,
                            original_metric_header=original_metric_header,
                            date_cell_value=date_value,
                            raw_value_text=None if value is None else str(value),
                        )
                    )
        return parsed

    def normalize(
        self,
        records: list[DphParsedRow],
        *,
        retrieved_at: datetime,
        raw_snapshot_id: int | None = None,
    ) -> list[ObservationRecord]:
        del records, retrieved_at, raw_snapshot_id
        raise SourceSchemaError(
            "Mass. DPH metric mapping is disabled until a real workbook is reviewed"
        )


def _header_map(
    raw_headers: tuple[object, ...], sheet_name: str
) -> dict[str, tuple[int, str]]:
    headers: dict[str, tuple[int, str]] = {}
    for index, value in enumerate(raw_headers):
        if value is None:
            continue
        if not isinstance(value, str):
            raise SourceSchemaError(f"sheet {sheet_name} contains a non-text header")
        normalized = _normalize_header(value)
        if normalized in headers:
            raise SourceSchemaError(
                f"sheet {sheet_name} contains duplicate normalized header {value}"
            )
        headers[normalized] = (index, value)
    return headers


def _normalize_header(value: str) -> str:
    return " ".join(value.split()).casefold()
